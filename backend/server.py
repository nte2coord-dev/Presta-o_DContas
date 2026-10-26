"""Sistema de Prestação de Contas — SEE/MG. FastAPI entrypoint."""
from dotenv import load_dotenv
from pathlib import Path

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / ".env")

import os
import uuid
import logging
from io import BytesIO
from datetime import datetime, timezone, date
from typing import List, Optional, Literal

from fastapi import FastAPI, APIRouter, HTTPException, Depends, Request, Response, UploadFile, File
from fastapi.responses import StreamingResponse
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field, EmailStr
import openpyxl

from auth import (
    hash_password, verify_password, create_access_token, create_refresh_token,
    set_auth_cookies, clear_auth_cookies, get_current_user, require_roles,
)
from docgen import gerar_parecer, gerar_capa, gerar_etiqueta, gerar_oficio

# --- DB ---
mongo_url = os.environ["MONGO_URL"]
client = AsyncIOMotorClient(mongo_url)
db = client[os.environ["DB_NAME"]]

app = FastAPI(title="Sistema de Prestação de Contas — SEE/MG")
api = APIRouter(prefix="/api")

# --- Helpers ---
def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()

def new_id() -> str:
    return str(uuid.uuid4())

def clean(doc: dict) -> dict:
    if doc:
        doc.pop("_id", None)
    return doc

# ========================= MODELS =========================
class LoginIn(BaseModel):
    email: EmailStr
    password: str

class UserCreate(BaseModel):
    email: EmailStr
    password: str
    name: str
    role: Literal["admin", "analista", "gestor"] = "analista"

class UserOut(BaseModel):
    id: str
    email: str
    name: str
    role: str

class EscolaIn(BaseModel):
    nome: str
    cnpj: str
    sre: str = "Metropolitana A"
    diretor: Optional[str] = ""
    telefone: Optional[str] = ""
    email: Optional[str] = ""
    endereco: Optional[str] = ""
    observacoes: Optional[str] = ""

class TermoIn(BaseModel):
    tipo: Literal["compromisso", "obras"]
    numero: str
    escola_id: str
    exercicio: int
    vigencia_inicio: str  # ISO date
    vigencia_fim: str
    valor_recebido: float = 0.0
    plano_trabalho: Optional[str] = ""
    analista_id: Optional[str] = None
    status: str = "EM_EXECUCAO"
    observacoes: Optional[str] = ""

class StatusUpdate(BaseModel):
    status: str
    observacoes: Optional[str] = ""

class AgendamentoIn(BaseModel):
    escola_id: str
    termo_id: Optional[str] = None
    data: str  # ISO datetime
    tipo_processo: Literal["compromisso", "obras", "outros"] = "compromisso"
    analista_id: Optional[str] = None
    observacoes: Optional[str] = ""

class OficioIn(BaseModel):
    numero: str
    termo_id: Optional[str] = None
    escola_id: str
    tipo: Literal["diligencia", "reiteracao", "outros"] = "diligencia"
    data_emissao: str
    prazo_dias: int = 30
    pendencias: List[str] = []
    status: Literal["aberto", "respondido", "vencido", "arquivado"] = "aberto"
    nlc: Optional[str] = ""
    observacoes: Optional[str] = ""

class MetaIn(BaseModel):
    analista_id: str
    exercicio: int
    processos_meta: int = 0
    processos_concluidos: int = 0
    observacoes: Optional[str] = ""

class LancamentoIn(BaseModel):
    termo_id: str
    tipo: Literal["aporte", "despesa", "rendimento", "devolucao"]
    data: str
    descricao: str
    valor: float
    fornecedor: Optional[str] = ""
    nota_fiscal: Optional[str] = ""

# ========================= STARTUP =========================
@app.on_event("startup")
async def on_startup():
    await db.users.create_index("email", unique=True)
    await db.users.create_index("id", unique=True)
    await db.escolas.create_index("id", unique=True)
    await db.escolas.create_index("cnpj")
    await db.termos.create_index("id", unique=True)
    await db.termos.create_index([("escola_id", 1), ("tipo", 1)])
    await db.agendamentos.create_index("id", unique=True)
    await db.oficios.create_index("id", unique=True)
    await db.metas.create_index("id", unique=True)
    await db.lancamentos.create_index("id", unique=True)
    await db.lancamentos.create_index("termo_id")
    # Seed admin
    admin_email = os.environ.get("ADMIN_EMAIL", "admin@seemg.gov.br")
    admin_password = os.environ.get("ADMIN_PASSWORD", "admin123")
    existing = await db.users.find_one({"email": admin_email})
    if not existing:
        await db.users.insert_one({
            "id": new_id(), "email": admin_email,
            "password_hash": hash_password(admin_password),
            "name": "Administrador", "role": "admin",
            "created_at": now_iso(),
        })
    elif not verify_password(admin_password, existing["password_hash"]):
        await db.users.update_one(
            {"email": admin_email},
            {"$set": {"password_hash": hash_password(admin_password)}}
        )
    # Seed demo analista + gestor (idempotent)
    for e, p, n, r in [
        ("analista@seemg.gov.br", "Analista@2026", "Analista Demo", "analista"),
        ("gestor@seemg.gov.br",   "Gestor@2026",   "Gestor Demo",   "gestor"),
    ]:
        if not await db.users.find_one({"email": e}):
            await db.users.insert_one({
                "id": new_id(), "email": e,
                "password_hash": hash_password(p),
                "name": n, "role": r, "created_at": now_iso(),
            })

@app.on_event("shutdown")
async def shutdown():
    client.close()

# ========================= AUTH ROUTES =========================
@api.post("/auth/login")
async def login(payload: LoginIn, response: Response):
    user = await db.users.find_one({"email": payload.email.lower()})
    if not user or not verify_password(payload.password, user["password_hash"]):
        raise HTTPException(401, "Credenciais inválidas")
    access = create_access_token(user["id"], user["email"], user["role"])
    refresh = create_refresh_token(user["id"])
    set_auth_cookies(response, access, refresh)
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"]}

@api.post("/auth/logout")
async def logout(response: Response, _user: dict = Depends(get_current_user)):
    clear_auth_cookies(response)
    return {"ok": True}

@api.get("/auth/me")
async def me(user: dict = Depends(get_current_user)):
    return {"id": user["id"], "email": user["email"], "name": user["name"], "role": user["role"]}

# ========================= USERS (admin) =========================
@api.get("/users")
async def list_users(_: dict = Depends(require_roles("admin", "gestor"))):
    docs = await db.users.find({}, {"password_hash": 0, "_id": 0}).to_list(500)
    return docs

@api.post("/users")
async def create_user(payload: UserCreate, _: dict = Depends(require_roles("admin"))):
    email = payload.email.lower()
    if await db.users.find_one({"email": email}):
        raise HTTPException(400, "E-mail já cadastrado")
    doc = {
        "id": new_id(), "email": email, "name": payload.name, "role": payload.role,
        "password_hash": hash_password(payload.password), "created_at": now_iso(),
    }
    await db.users.insert_one(doc)
    doc.pop("password_hash"); doc.pop("_id", None)
    return doc

@api.delete("/users/{uid}")
async def delete_user(uid: str, _: dict = Depends(require_roles("admin"))):
    await db.users.delete_one({"id": uid})
    return {"ok": True}

# ========================= ESCOLAS =========================
@api.get("/escolas")
async def list_escolas(_: dict = Depends(get_current_user)):
    return await db.escolas.find({}, {"_id": 0}).to_list(1000)

@api.post("/escolas")
async def create_escola(payload: EscolaIn, _: dict = Depends(require_roles("admin", "analista"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.escolas.insert_one(doc.copy())
    return clean(doc)

@api.put("/escolas/{eid}")
async def update_escola(eid: str, payload: EscolaIn, _: dict = Depends(require_roles("admin", "analista"))):
    await db.escolas.update_one({"id": eid}, {"$set": payload.model_dump()})
    return clean(await db.escolas.find_one({"id": eid}))

@api.delete("/escolas/{eid}")
async def delete_escola(eid: str, _: dict = Depends(require_roles("admin"))):
    await db.escolas.delete_one({"id": eid})
    return {"ok": True}

# ========================= TERMOS =========================
@api.get("/termos")
async def list_termos(tipo: Optional[str] = None, _: dict = Depends(get_current_user)):
    q = {}
    if tipo:
        q["tipo"] = tipo
    return await db.termos.find(q, {"_id": 0}).to_list(2000)

@api.post("/termos")
async def create_termo(payload: TermoIn, _: dict = Depends(require_roles("admin", "analista"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.termos.insert_one(doc.copy())
    return clean(doc)

@api.put("/termos/{tid}")
async def update_termo(tid: str, payload: TermoIn, _: dict = Depends(require_roles("admin", "analista"))):
    await db.termos.update_one({"id": tid}, {"$set": payload.model_dump()})
    return clean(await db.termos.find_one({"id": tid}))

@api.patch("/termos/{tid}/status")
async def update_status(tid: str, payload: StatusUpdate, _: dict = Depends(require_roles("admin", "analista"))):
    await db.termos.update_one({"id": tid}, {"$set": {"status": payload.status, "observacoes_status": payload.observacoes}})
    return clean(await db.termos.find_one({"id": tid}))

@api.delete("/termos/{tid}")
async def delete_termo(tid: str, _: dict = Depends(require_roles("admin"))):
    await db.termos.delete_one({"id": tid})
    await db.lancamentos.delete_many({"termo_id": tid})
    return {"ok": True}

# ========================= AGENDAMENTOS =========================
@api.get("/agendamentos")
async def list_ag(_: dict = Depends(get_current_user)):
    return await db.agendamentos.find({}, {"_id": 0}).sort("data", 1).to_list(2000)

@api.post("/agendamentos")
async def create_ag(payload: AgendamentoIn, _: dict = Depends(require_roles("admin", "analista"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.agendamentos.insert_one(doc.copy())
    return clean(doc)

@api.put("/agendamentos/{aid}")
async def update_ag(aid: str, payload: AgendamentoIn, _: dict = Depends(require_roles("admin", "analista"))):
    await db.agendamentos.update_one({"id": aid}, {"$set": payload.model_dump()})
    return clean(await db.agendamentos.find_one({"id": aid}))

@api.delete("/agendamentos/{aid}")
async def delete_ag(aid: str, _: dict = Depends(require_roles("admin", "analista"))):
    await db.agendamentos.delete_one({"id": aid})
    return {"ok": True}

# ========================= OFÍCIOS =========================
@api.get("/oficios")
async def list_of(_: dict = Depends(get_current_user)):
    return await db.oficios.find({}, {"_id": 0}).sort("data_emissao", -1).to_list(2000)

@api.post("/oficios")
async def create_of(payload: OficioIn, _: dict = Depends(require_roles("admin", "analista"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.oficios.insert_one(doc.copy())
    return clean(doc)

@api.put("/oficios/{oid}")
async def update_of(oid: str, payload: OficioIn, _: dict = Depends(require_roles("admin", "analista"))):
    await db.oficios.update_one({"id": oid}, {"$set": payload.model_dump()})
    return clean(await db.oficios.find_one({"id": oid}))

@api.delete("/oficios/{oid}")
async def delete_of(oid: str, _: dict = Depends(require_roles("admin"))):
    await db.oficios.delete_one({"id": oid})
    return {"ok": True}

# ========================= METAS =========================
@api.get("/metas")
async def list_metas(_: dict = Depends(get_current_user)):
    return await db.metas.find({}, {"_id": 0}).to_list(1000)

@api.post("/metas")
async def create_meta(payload: MetaIn, _: dict = Depends(require_roles("admin", "gestor"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.metas.insert_one(doc.copy())
    return clean(doc)

@api.put("/metas/{mid}")
async def update_meta(mid: str, payload: MetaIn, _: dict = Depends(require_roles("admin", "gestor"))):
    await db.metas.update_one({"id": mid}, {"$set": payload.model_dump()})
    return clean(await db.metas.find_one({"id": mid}))

@api.delete("/metas/{mid}")
async def delete_meta(mid: str, _: dict = Depends(require_roles("admin"))):
    await db.metas.delete_one({"id": mid})
    return {"ok": True}

# ========================= CALCULADORA / LANÇAMENTOS =========================
@api.get("/lancamentos")
async def list_lanc(termo_id: Optional[str] = None, _: dict = Depends(get_current_user)):
    q = {"termo_id": termo_id} if termo_id else {}
    return await db.lancamentos.find(q, {"_id": 0}).sort("data", 1).to_list(5000)

@api.post("/lancamentos")
async def create_lanc(payload: LancamentoIn, _: dict = Depends(require_roles("admin", "analista"))):
    doc = {"id": new_id(), **payload.model_dump(), "created_at": now_iso()}
    await db.lancamentos.insert_one(doc.copy())
    return clean(doc)

@api.delete("/lancamentos/{lid}")
async def delete_lanc(lid: str, _: dict = Depends(require_roles("admin", "analista"))):
    await db.lancamentos.delete_one({"id": lid})
    return {"ok": True}

@api.get("/lancamentos/{termo_id}/resumo")
async def resumo_termo(termo_id: str, _: dict = Depends(get_current_user)):
    docs = await db.lancamentos.find({"termo_id": termo_id}, {"_id": 0}).to_list(5000)
    aportes = sum(d["valor"] for d in docs if d["tipo"] == "aporte")
    despesas = sum(d["valor"] for d in docs if d["tipo"] == "despesa")
    rendimentos = sum(d["valor"] for d in docs if d["tipo"] == "rendimento")
    devolucao = sum(d["valor"] for d in docs if d["tipo"] == "devolucao")
    saldo = aportes + rendimentos - despesas - devolucao
    return {
        "aportes": aportes, "despesas": despesas, "rendimentos": rendimentos,
        "devolucao": devolucao, "saldo": saldo, "qtd": len(docs),
    }

# ========================= DASHBOARD =========================
@api.get("/dashboard/stats")
async def dashboard(_: dict = Depends(get_current_user)):
    termos = await db.termos.find({}, {"_id": 0}).to_list(5000)
    by_status: dict[str, int] = {}
    for t in termos:
        by_status[t["status"]] = by_status.get(t["status"], 0) + 1
    oficios = await db.oficios.find({}, {"_id": 0}).to_list(5000)
    abertos = [o for o in oficios if o.get("status") == "aberto"]
    hoje = datetime.now(timezone.utc).date()
    criticos = 0
    for o in abertos:
        try:
            emissao = datetime.fromisoformat(o["data_emissao"].replace("Z", "+00:00")).date()
            dias_rest = o.get("prazo_dias", 30) - (hoje - emissao).days
            if dias_rest <= 5:
                criticos += 1
        except Exception:
            pass
    return {
        "escolas": await db.escolas.count_documents({}),
        "termos_total": len(termos),
        "termos_compromisso": sum(1 for t in termos if t.get("tipo") == "compromisso"),
        "termos_obras": sum(1 for t in termos if t.get("tipo") == "obras"),
        "em_analise": by_status.get("EM_ANALISE", 0),
        "diligencia_emitida": by_status.get("DILIGENCIA_EMITIDA", 0),
        "aprovados": by_status.get("APROVADO_SEM_RESSALVA", 0) + by_status.get("APROVADO_COM_RESSALVA", 0),
        "bloqueados_siafi": by_status.get("BLOQUEADO_SIAFI", 0),
        "arquivados": by_status.get("ARQUIVADO", 0),
        "oficios_abertos": len(abertos),
        "oficios_criticos": criticos,
        "status_distribution": by_status,
    }

# ========================= XLSX IMPORT / EXPORT =========================
def _xlsx_response(wb, filename: str) -> StreamingResponse:
    buf = BytesIO()
    wb.save(buf); buf.seek(0)
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})

@api.get("/export/{resource}")
async def export_xlsx(resource: str, _: dict = Depends(get_current_user)):
    allowed = {"escolas", "termos", "oficios", "agendamentos", "metas", "lancamentos"}
    if resource not in allowed:
        raise HTTPException(400, "Recurso inválido")
    data = await getattr(db, resource).find({}, {"_id": 0}).to_list(10000)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = resource.capitalize()
    if not data:
        ws.append(["(sem dados)"])
    else:
        headers = sorted({k for d in data for k in d.keys()})
        ws.append(headers)
        for d in data:
            ws.append([str(d.get(h, "")) for h in headers])
    return _xlsx_response(wb, f"{resource}_{datetime.now().strftime('%Y%m%d')}.xlsx")

@api.post("/import/{resource}")
async def import_xlsx(resource: str, file: UploadFile = File(...),
                     _: dict = Depends(require_roles("admin", "analista"))):
    allowed = {"escolas", "oficios", "agendamentos", "metas"}
    if resource not in allowed:
        raise HTTPException(400, "Recurso não suportado para importação")
    content = await file.read()
    wb = openpyxl.load_workbook(BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return {"inserted": 0}
    headers = [str(h or "").strip() for h in rows[0]]
    inserted = 0
    for r in rows[1:]:
        obj = {h: (v if v is not None else "") for h, v in zip(headers, r) if h}
        if not any(obj.values()):
            continue
        obj["id"] = new_id()
        obj["created_at"] = now_iso()
        await getattr(db, resource).insert_one(obj)
        inserted += 1
    return {"inserted": inserted}

# ========================= GERAÇÃO DE DOCUMENTOS =========================
class DocRequest(BaseModel):
    kind: Literal["parecer", "capa", "etiqueta", "oficio"]
    data: dict

@api.post("/documentos/gerar")
async def gerar_doc(req: DocRequest, _: dict = Depends(get_current_user)):
    try:
        if req.kind == "parecer":
            buf = gerar_parecer(req.data)
            fname = "parecer_conclusivo.docx"
        elif req.kind == "capa":
            buf = gerar_capa(req.data)
            fname = "capa_processo.docx"
        elif req.kind == "etiqueta":
            buf = gerar_etiqueta(req.data)
            fname = "etiqueta_caixa.docx"
        else:
            buf = gerar_oficio(req.data)
            fname = "oficio_diligencia.docx"
    except Exception as e:
        raise HTTPException(400, f"Erro ao gerar documento: {e}")
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'})

@api.get("/")
async def root():
    return {"service": "Prestação de Contas SEE/MG", "status": "ok"}

# --- Register & CORS ---
app.include_router(api)

frontend_url = os.environ.get("FRONTEND_URL", "http://localhost:3000")
cors_origins = os.environ.get("CORS_ORIGINS", frontend_url).split(",")
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins if cors_origins != ["*"] else ["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(name)s - %(levelname)s - %(message)s")
logger = logging.getLogger(__name__)
