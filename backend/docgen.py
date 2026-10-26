"""Document generation helpers (DOCX) for Parecer, Capa, Etiquetas and Ofícios."""
from io import BytesIO
from datetime import datetime
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


def _set_heading(doc, text, level=1, center=True):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(16 if level == 1 else 13)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def _add_para(doc, text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    p.alignment = align
    return p


def _br(doc):
    doc.add_paragraph()


def _buf(doc) -> BytesIO:
    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


def gerar_parecer(data: dict) -> BytesIO:
    """
    data: {
      tipo: 'sem_ressalva' | 'com_ressalva',
      escola, cnpj, termo_numero, exercicio, analista,
      valor_recebido, valor_aplicado, saldo_devolver, rendimentos,
      ressalvas (list), observacoes
    }
    """
    doc = Document()
    _set_heading(doc, "SECRETARIA DE ESTADO DE EDUCAÇÃO DE MINAS GERAIS", level=1)
    _set_heading(doc, "SUPERINTENDÊNCIA REGIONAL DE ENSINO METROPOLITANA A", level=2)
    _br(doc)
    tipo = (data.get("tipo") or "sem_ressalva").upper()
    titulo = "PARECER CONCLUSIVO SEM RESSALVAS" if tipo == "SEM_RESSALVA" else "PARECER CONCLUSIVO COM RESSALVAS"
    _set_heading(doc, titulo, level=2)
    _br(doc)
    _add_para(doc, f"Caixa Escolar: {data.get('escola','—')}", bold=True)
    _add_para(doc, f"CNPJ: {data.get('cnpj','—')}")
    _add_para(doc, f"Termo nº: {data.get('termo_numero','—')} — Exercício: {data.get('exercicio','—')}")
    _add_para(doc, f"Analista Responsável: {data.get('analista','—')}")
    _br(doc)
    _add_para(doc, "1. DOS RECURSOS", bold=True)
    _add_para(doc, f"Valor recebido: R$ {data.get('valor_recebido','0,00')}")
    _add_para(doc, f"Valor aplicado: R$ {data.get('valor_aplicado','0,00')}")
    _add_para(doc, f"Rendimentos: R$ {data.get('rendimentos','0,00')}")
    _add_para(doc, f"Saldo a devolver: R$ {data.get('saldo_devolver','0,00')}")
    _br(doc)
    if tipo == "COM_RESSALVA":
        _add_para(doc, "2. DAS RESSALVAS", bold=True)
        for r in (data.get("ressalvas") or []):
            _add_para(doc, f"• {r}")
        _br(doc)
    _add_para(doc, "3. CONCLUSÃO", bold=True)
    base = ("Após análise documental e financeira dos recursos repassados à caixa escolar, "
            "verificou-se a adequada aplicação dos valores em consonância com o plano de trabalho pactuado.")
    if tipo == "COM_RESSALVA":
        base += " As ressalvas acima apontadas foram devidamente registradas e justificadas."
    _add_para(doc, base)
    _add_para(doc, data.get("observacoes", ""))
    _br(doc); _br(doc)
    _add_para(doc, f"Belo Horizonte, {datetime.now().strftime('%d/%m/%Y')}", align=WD_ALIGN_PARAGRAPH.RIGHT)
    _br(doc); _br(doc)
    _add_para(doc, "_______________________________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, f"{data.get('analista','Analista')}", align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, "Analista Técnico(a) — Prestação de Contas", align=WD_ALIGN_PARAGRAPH.CENTER)
    return _buf(doc)


def gerar_capa(data: dict) -> BytesIO:
    doc = Document()
    _set_heading(doc, "SECRETARIA DE ESTADO DE EDUCAÇÃO DE MINAS GERAIS", level=1)
    _set_heading(doc, "SRE METROPOLITANA A — PRESTAÇÃO DE CONTAS", level=2)
    _br(doc); _br(doc)
    _set_heading(doc, "CAPA DE PROCESSO", level=1)
    _br(doc); _br(doc)
    for label, key in [("Caixa Escolar", "escola"), ("CNPJ", "cnpj"),
                       ("Termo nº", "termo_numero"), ("Exercício", "exercicio"),
                       ("Tipo", "tipo_processo"), ("Analista", "analista")]:
        _add_para(doc, f"{label}: {data.get(key,'—')}", size=13)
    _br(doc)
    _add_para(doc, f"Volume: {data.get('volume','01')} — Folhas: {data.get('folhas','—')}", size=13)
    _br(doc); _br(doc); _br(doc)
    _add_para(doc, f"Belo Horizonte, {datetime.now().strftime('%d/%m/%Y')}",
              align=WD_ALIGN_PARAGRAPH.RIGHT, size=11)
    return _buf(doc)


def gerar_etiqueta(data: dict) -> BytesIO:
    doc = Document()
    _set_heading(doc, "ETIQUETA — CAIXA BOX DE PRESTAÇÃO DE CONTAS", level=1)
    _br(doc)
    _add_para(doc, f"SRE: {data.get('sre','Metropolitana A')}", bold=True, size=14)
    _add_para(doc, f"Caixa Escolar: {data.get('escola','—')}", bold=True, size=14)
    _add_para(doc, f"CNPJ: {data.get('cnpj','—')}", size=12)
    _add_para(doc, f"Exercício(s): {data.get('exercicio','—')}", size=12)
    _add_para(doc, f"Tipo: {data.get('tipo_processo','—')}", size=12)
    _add_para(doc, f"Nº Caixa: {data.get('numero_caixa','—')}", size=12)
    _add_para(doc, f"Localização: {data.get('localizacao','—')}", size=12)
    return _buf(doc)


def gerar_oficio(data: dict) -> BytesIO:
    doc = Document()
    _set_heading(doc, "SECRETARIA DE ESTADO DE EDUCAÇÃO DE MINAS GERAIS", level=1)
    _set_heading(doc, "SRE METROPOLITANA A — DIVISÃO DE PRESTAÇÃO DE CONTAS", level=2)
    _br(doc)
    _add_para(doc, f"Ofício nº {data.get('numero','—')} / {datetime.now().year}", bold=True)
    _add_para(doc, f"Belo Horizonte, {datetime.now().strftime('%d/%m/%Y')}",
              align=WD_ALIGN_PARAGRAPH.RIGHT)
    _br(doc)
    _add_para(doc, f"À Caixa Escolar: {data.get('escola','—')}", bold=True)
    _add_para(doc, f"Assunto: Diligência — Termo nº {data.get('termo_numero','—')}")
    _br(doc)
    _add_para(doc, "Prezados(as),")
    _br(doc)
    _add_para(doc,
        f"Em análise à prestação de contas referente ao Termo nº {data.get('termo_numero','—')}, "
        f"exercício {data.get('exercicio','—')}, foram identificadas as pendências abaixo, "
        f"cuja regularização deverá ocorrer no prazo de {data.get('prazo_dias',30)} dias, "
        "contados do recebimento deste ofício, sob pena de bloqueio no SIAFI e demais medidas cabíveis.")
    _br(doc)
    _add_para(doc, "PENDÊNCIAS:", bold=True)
    for p in (data.get("pendencias") or []):
        _add_para(doc, f"• {p}")
    _br(doc)
    _add_para(doc, data.get("observacoes", ""))
    _br(doc); _br(doc)
    _add_para(doc, "Atenciosamente,", align=WD_ALIGN_PARAGRAPH.CENTER)
    _br(doc)
    _add_para(doc, "_______________________________________", align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, f"{data.get('analista','Analista')}", align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, "Analista Técnico(a) — Prestação de Contas", align=WD_ALIGN_PARAGRAPH.CENTER)
    return _buf(doc)
